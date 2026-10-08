import sys
import os
import json
import shutil
import stat

def process_word_document(template_path, output_path, data):
    """
    Process Word template (.doc or .docx), perform text replacements and table updates.
    Preserves original template file completely (R17).
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found: {template_path}")

    output_dir = os.path.dirname(output_path)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir, exist_ok=True)

    # Ensure output file is not read-only
    if os.path.exists(output_path):
        try:
            os.chmod(output_path, stat.S_IWRITE | stat.S_IREAD)
        except Exception:
            pass

    if not os.path.exists(output_path) or os.path.abspath(template_path) != os.path.abspath(output_path):
        shutil.copy2(template_path, output_path)

    try:
        os.chmod(output_path, stat.S_IWRITE | stat.S_IREAD)
    except Exception:
        pass

    doc_type = data.get('type', 'cert') # 'cert' or 'packing'
    form_data = data.get('formData', {})

    is_windows = sys.platform.startswith('win')

    if is_windows:
        try:
            import win32com.client
            app = None
            try:
                app = win32com.client.Dispatch("KWps.Application")
            except Exception:
                try:
                    app = win32com.client.Dispatch("Wps.Application")
                except Exception:
                    app = win32com.client.Dispatch("Word.Application")

            if app is not None:
                app.Visible = False
                try:
                    app.DisplayAlerts = 0
                except Exception:
                    pass

                doc = app.Documents.Open(os.path.abspath(output_path))

                model = str(form_data.get('model', ''))
                device_sn = str(form_data.get('deviceSn', ''))
                shipping_location = str(form_data.get('shippingLocation') or '')
                sensor_model = str(form_data.get('sensorModel', ''))
                sensor_sn = str(form_data.get('sensorSn', ''))
                cert_date = str(form_data.get('certDate') or form_data.get('date') or '')
                has_pump = bool(form_data.get('hasPump', True))
                # 保护行/传感器行由该任务绑定模板的 protectedRows/packingItems 决定，
                # 不再用型号名称（含 POA）推导 (整改 B.1/B.5)，否则 POA3500 会被误当成 POA200
                is_pump_model = 'POA' in model.upper()
                field_mappings = data.get('fieldMappings') or data.get('field_mappings') or {}
                packing_cfg_items = field_mappings.get('packingItems') or []
                has_sensor_row_configured = False
                sensor_row_name = '传感器'
                sensor_data_index = -1
                for _idx, _cfg in enumerate(packing_cfg_items):
                    if _cfg.get('isProtectedSensor') or str(_cfg.get('name') or '') == '传感器':
                        has_sensor_row_configured = True
                        sensor_data_index = _idx
                        if _cfg.get('name'):
                            sensor_row_name = str(_cfg.get('name'))
                        break
                if not has_sensor_row_configured and sensor_model:
                    has_sensor_row_configured = True

                def set_cell_text(tbl, row, col, value):
                    """完整覆写单元格文本，避免 Range.Text 在存在格式标记时变成追加（如 19.998 mA mA）"""
                    if value is None:
                        return False
                    try:
                        rng = tbl.Cell(row, col).Range
                        try:
                            rng.MoveEnd(1, -1)  # wdCharacter=1，排除单元格结束标记
                        except Exception:
                            pass
                        rng.Text = str(value)
                        return True
                    except Exception:
                        return False

                # Replace default SN AP10007513 in StoryRanges
                if device_sn and device_sn != "AP10007513":
                    for story in doc.StoryRanges:
                        try:
                            find = story.Find
                            find.Text = "AP10007513"
                            find.Replacement.Text = device_sn
                            find.Execute(Replace=2)
                        except Exception:
                            pass

                if doc.Tables.Count >= 1:
                    table = doc.Tables.Item(1)

                    if doc_type == 'cert':
                        ambient_temp = str(form_data.get('ambientTemp', '')) if form_data.get('ambientTemp') is not None else ''
                        relative_humidity = str(form_data.get('relativeHumidity', '')) if form_data.get('relativeHumidity') is not None else ''

                        if cert_date and table.Rows.Count >= 4 and table.Columns.Count >= 6:
                            table.Cell(4, 6).Range.Text = cert_date

                        if device_sn and table.Rows.Count >= 7 and table.Columns.Count >= 2:
                            table.Cell(7, 2).Range.Text = device_sn

                        test_points = form_data.get('testPoints', [])
                        field_mappings = data.get('fieldMappings') or data.get('field_mappings') or {}
                        if not field_mappings and cert_template:
                            field_mappings = cert_template.get('field_mappings') or {}
                        tc = field_mappings.get('tableConfig') or data.get('tableConfig') or (cert_template.get('field_mappings', {}).get('tableConfig') if cert_template else None)
                        columns = tc.get('columns', []) if tc else []
                        start_r = tc.get('startRow', 12) + 1 if tc and tc.get('startRow') is not None else 13

                        needed_rows = start_r + len(test_points) - 1
                        while table.Rows.Count < needed_rows:
                            try: table.Rows.Add()
                            except: break
                        while table.Rows.Count > needed_rows and table.Rows.Count > start_r:
                            try: table.Rows.Item(table.Rows.Count).Delete()
                            except: break

                        if columns and start_r > 1:
                            header_r = tc.get('headerRow', start_r - 2) + 1 if 'headerRow' in tc else start_r - 1
                            if 1 <= header_r <= table.Rows.Count:
                                for col_def in columns:
                                    c_idx = col_def.get('colIdx', 0) + 1
                                    c_lbl = col_def.get('label')
                                    if c_lbl:
                                        try:
                                            table.Cell(header_r, c_idx).Range.Text = str(c_lbl)
                                        except Exception:
                                            try: table.Rows.Item(header_r).Cells.Item(c_idx).Range.Text = str(c_lbl)
                                            except Exception: pass

                        for idx, tp in enumerate(test_points):
                            r_idx = start_r + idx
                            if table.Rows.Count >= r_idx:
                                if columns:
                                    for col_def in columns:
                                        c_idx = col_def.get('colIdx', 0) + 1
                                        c_key = col_def.get('key')
                                        val = None
                                        vals = tp.get('values', {})
                                        if c_key in vals: val = vals[c_key]
                                        elif str(col_def.get('colIdx')) in vals: val = vals[str(col_def.get('colIdx'))]
                                        elif col_def.get('label') in vals: val = vals[col_def.get('label')]
                                        elif col_def.get('isSeq') and tp.get('name'): val = tp.get('name')
                                        elif col_def.get('isStd') and tp.get('std') is not None: val = tp.get('std')
                                        elif col_def.get('isAct') and tp.get('act') is not None: val = tp.get('act')
                                        if val is not None:
                                            try:
                                                table.Cell(r_idx, c_idx).Range.Text = str(val)
                                            except Exception:
                                                try: table.Rows.Item(r_idx).Cells.Item(c_idx).Range.Text = str(val)
                                                except Exception: pass
                                else:
                                    std_val = tp.get('std') if tp.get('std') is not None else tp.get('standard')
                                    act_val = tp.get('act') if tp.get('act') is not None else tp.get('actual')
                                    if std_val is not None and table.Columns.Count >= 2:
                                        table.Cell(r_idx, 2).Range.Text = str(std_val)
                                    if act_val is not None and table.Columns.Count >= 3:
                                        table.Cell(r_idx, 3).Range.Text = str(act_val)
                    elif doc_type == 'packing':
                        pump_str = "带泵" if (is_pump_model and has_pump) else ""
                        main_remark = f"SN: {device_sn} {pump_str}".strip() if pump_str else f"SN: {device_sn}"
                        sensor_remark_str = f"SN: {sensor_sn}" if sensor_sn else ""
                        packing_items = form_data.get('packingItems', [])
                        first_data_row = 2

                        # 按行名称识别真实存在的保护行，不假设“第二行就是传感器行” (整改 B.5/B.6)
                        for _r in range(2, table.Rows.Count + 1):
                            try:
                                _name = str(table.Cell(_r, 2).Range.Text).strip()
                            except Exception:
                                continue
                            if _name.startswith('主设备'):
                                if model:
                                    set_cell_text(table, _r, 3, model)
                                set_cell_text(table, _r, 7, main_remark)
                            elif has_sensor_row_configured and _name == sensor_row_name:
                                if sensor_model:
                                    set_cell_text(table, _r, 3, sensor_model)
                                set_cell_text(table, _r, 7, sensor_remark_str)

                        if packing_items:
                            # 若模板配置未指明传感器行，则按条目名称/标记推断（仅当模板确实配置了传感器行时才生效）
                            if sensor_data_index < 0 and has_sensor_row_configured:
                                for _i, _it in enumerate(packing_items):
                                    if str(_it.get('name') or '') == sensor_row_name or _it.get('isProtectedSensor'):
                                        sensor_data_index = _i
                                        break

                            # 只在数据区增删行：在最后一行数据行的某个单元格上 Range.Rows.Add()，
                            # 新行会插入到该数据行之后、表格下方内容之前 (整改 C.6)
                            template_data_rows = max(0, table.Rows.Count - first_data_row + 1)
                            current_data_rows = template_data_rows
                            guard = 0
                            while current_data_rows < len(packing_items) and guard < 60:
                                guard += 1
                                added = False
                                last_row = first_data_row + current_data_rows - 1
                                for _c in range(table.Columns.Count, 0, -1):
                                    try:
                                        table.Cell(last_row, _c).Range.Rows.Add()
                                        added = True
                                        break
                                    except Exception:
                                        continue
                                if not added:
                                    break
                                current_data_rows += 1

                            for idx, item in enumerate(packing_items):
                                r = first_data_row + idx
                                if r < first_data_row or r > table.Rows.Count:
                                    break
                                set_cell_text(table, r, 1, str(idx + 1))
                                if item.get('name') is not None:
                                    set_cell_text(table, r, 2, str(item.get('name')))
                                # 规格：主设备用型号；只有真实传感器行才用传感器型号，其余保留用户填写值 (整改 B.6)
                                if idx == 0 and model:
                                    set_cell_text(table, r, 3, model)
                                elif sensor_data_index >= 0 and idx == sensor_data_index and sensor_model:
                                    set_cell_text(table, r, 3, sensor_model)
                                elif item.get('spec') is not None:
                                    set_cell_text(table, r, 3, str(item.get('spec')))
                                if item.get('count') is not None:
                                    set_cell_text(table, r, 4, str(item.get('count')))
                                if item.get('unit') is not None:
                                    set_cell_text(table, r, 5, str(item.get('unit')))
                                if item.get('standard') is not None:
                                    set_cell_text(table, r, 6, str(item.get('standard')))
                                # 备注：仅主设备行与真实传感器行使用自动备注 (整改 B.6)
                                if idx == 0:
                                    set_cell_text(table, r, 7, main_remark)
                                elif sensor_data_index >= 0 and idx == sensor_data_index:
                                    set_cell_text(table, r, 7, sensor_remark_str)
                                elif item.get('remark') is not None:
                                    set_cell_text(table, r, 7, str(item.get('remark')))

                doc.Save()
                doc.Close(False)
                app.Quit()
                return
        except Exception as e:
            print(f"win32com processing fallback: {e}", file=sys.stderr)

    # Cross-platform / python-docx fallback
    try:
        import docx
        if output_path.endswith('.docx'):
            doc = docx.Document(output_path)
            device_sn = form_data.get('deviceSn', '')
            if device_sn:
                for p in doc.paragraphs:
                    if 'AP10007513' in p.text:
                        p.text = p.text.replace('AP10007513', device_sn)
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if 'AP10007513' in cell.text:
                                cell.text = cell.text.replace('AP10007513', device_sn)
            doc.save(output_path)
    except Exception as e:
        print(f"docx processing notice: {e}", file=sys.stderr)

    print(f"Document processed successfully: {output_path}")

if __name__ == '__main__':
    if len(sys.argv) < 4:
        print("Usage: doc_processor.py <template_path> <output_path> <json_data_path>")
        sys.exit(1)

    template_p = sys.argv[1]
    output_p = sys.argv[2]
    json_data_p = sys.argv[3]

    with open(json_data_p, 'r', encoding='utf-8') as f:
        data_content = json.load(f)

    process_word_document(template_p, output_p, data_content)
